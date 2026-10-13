"""
seed_co.py  — populate course_outcomes table from CO-PO Mapping excel.
Usage:
    python seed_co.py --xlsx "CO-PO Mapping R22 .xlsx"
"""

import argparse
from openpyxl import load_workbook

from app import app
from models import db, Subject, CourseOutcome


def load_co_map(xlsx_path):
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    co_map = {}
    current_subject = None

    for row in ws.iter_rows(values_only=True):
        # pad short rows
        row = list(row) + [None] * 6
        col1, col2, col3, col4 = row[0], row[1], row[2], row[3]

        co3 = str(col3).strip() if col3 else ''
        co4 = str(col4).strip() if col4 else ''

        # new subject row: col1 has SL number, col2 has subject name
        if col2 and co3.startswith('CO') and len(co3) <= 4 and col4:
            subject_name = str(col2).strip()
            if subject_name.upper() != 'COURSES':
                current_subject = subject_name.upper().strip()
                co_map.setdefault(current_subject, [])
            if current_subject and co3 not in [x['co'] for x in co_map[current_subject]]:
                co_map[current_subject].append({'co': co3, 'statement': co4})

        # continuation row: col2 empty, col3 is CO
        elif not col2 and co3.startswith('CO') and len(co3) <= 4 and col4 and current_subject:
            if co3 not in [x['co'] for x in co_map[current_subject]]:
                co_map[current_subject].append({'co': co3, 'statement': co4})

    wb.close()
    return co_map


def normalize(name):
    return ' '.join(name.upper().strip().split())


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--xlsx', required=True)
    parser.add_argument('--dry-run', action='store_true')
    args = parser.parse_args()

    co_map = load_co_map(args.xlsx)
    print(f"Parsed {len(co_map)} subjects from xlsx\n")

    with app.app_context():
        all_subjects = Subject.query.all()
        db_map = {normalize(s.subject_name): s for s in all_subjects}

        inserted = 0
        skipped = 0
        missed = []

        for xl_name, cos in co_map.items():
            norm = normalize(xl_name)
            subject = db_map.get(norm)

            if subject is None:
                # fuzzy: one side contains the other
                candidates = [k for k in db_map if norm in k or k in norm]
                if len(candidates) == 1:
                    subject = db_map[candidates[0]]
                    print(f"  fuzzy: '{xl_name}' -> '{subject.subject_name}'")
                else:
                    missed.append(xl_name)
                    continue

            if args.dry_run:
                print(f"  [DRY] {subject.subject_name}: {len(cos)} COs")
                continue

            for co_entry in cos:
                existing = CourseOutcome.query.filter_by(
                    subject_id=subject.id, co_number=co_entry['co']
                ).first()
                if existing:
                    skipped += 1
                    continue
                db.session.add(CourseOutcome(
                    subject_id=subject.id,
                    co_number=co_entry['co'],
                    statement=co_entry['statement']
                ))
                inserted += 1

        if not args.dry_run:
            db.session.commit()

        print(f"\nInserted: {inserted}  Skipped (already exist): {skipped}")
        if missed:
            print(f"\nMISSED ({len(missed)} subjects not in DB):")
            for m in missed:
                print(f"  - {m}")


if __name__ == '__main__':
    main()