"""
fix_missed_cos.py — manually map missed xlsx subjects to DB subjects and import COs
"""
from app import app
from models import db, Subject, CourseOutcome
from openpyxl import load_workbook

MANUAL_MAP = {
    'APPLIED PHYSICS': 'Advanced Engineering Physics',
    'ELEMENTS OF COMPUTER SCIENCE AND ENGINEERING': None,
    'ENVIRONMENTA SCIENCE': 'Environmental Science',
    'OBJECT ORIENTED PROGRAMMING USING JAVA': 'Object Oriented Programming through Java',
    'SKILL DEVELOPMENT COURSE( NODE JS/REACT JS/DJANGO)': None,
    'CONSTITUTION OF INDIA': None,
    'HUMAN VALUES AND ETHICS': 'Human Value and Ethics',
    'MATHEMATICAL STATISTICAL FOUNDATIONS': 'Mathematical and Statistical Foundations',
    'DATABASE MANAGEMENT SYSTEMS': 'Database Management System',
    'SKILL DEVELOPMENT COURSE(PROLOG/ LISP/ PYSWIP)': None,
    'DESIGN AND ANALYSIS OF ALGORITHMS': 'Algorithm Design and Analysis',
    'BUSINESS ECONOMICS AND FINANCIAL ANALYSIS': 'Business Economics & Financial Analysis',
    'PE-1 INTODUCTION TO DATA SCIENCE': 'Introduction to Data Science',
    'ADVANCED ENGLISH COMMUNICATION SKILLS LAB': None,
    'LOGICAL REASONING -I': 'Logical Reasoning \u2013 I',
    'NATURAL LANGUAGE PROCESSING LAB/ FULL STACK DEVELOPMENT LAB': None,
    'NATURE INSPIRED COMPUTIUNG': 'Nature Inspired Computing',
    'PE-III LAB(GENAI LAB)': None,
    'PROJECT STAGE - I': None,
    'PE-IV (REINFORCEMENTLEARNING)': 'Reinforcement Learning',
    'OE-III ( ENVIRONMENTAL POLLUTION AND CONTROL)': None,
    'PROJECT STAGE - II': None,
    'IT WORKSHOP': None,
    'COMPUTER AIDED ENGINEERING GRAPHICS': 'Engineering Drawing and Computer Aided Drafting',
    'ENGLISH LANGAUGE COMMUNICATION SKILLS LAB': None,
    'APPLIES PHYSICS LAB': None,
    'ENGINEERING WORKSHOP': None,
    'GENDER SENSITIZATION LAB': None,
}

def load_co_map(xlsx_path):
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    co_map = {}
    current_subject = None
    for row in ws.iter_rows(values_only=True):
        row = list(row) + [None]*6
        col1, col2, col3, col4 = row[0], row[1], row[2], row[3]
        co3 = str(col3).strip() if col3 else ''
        co4 = str(col4).strip() if col4 else ''
        if col2 and co3.startswith('CO') and len(co3) <= 4 and col4:
            subject_name = str(col2).strip()
            if subject_name.upper() != 'COURSES':
                current_subject = subject_name.upper().strip()
                co_map.setdefault(current_subject, [])
            if current_subject and co3 not in [x['co'] for x in co_map[current_subject]]:
                co_map[current_subject].append({'co': co3, 'statement': co4})
        elif not col2 and co3.startswith('CO') and len(co3) <= 4 and col4 and current_subject:
            if co3 not in [x['co'] for x in co_map[current_subject]]:
                co_map[current_subject].append({'co': co3, 'statement': co4})
    wb.close()
    return co_map

xlsx = 'CO-PO Mapping R22 .xlsx'
co_map = load_co_map(xlsx)

inserted = 0
skipped = 0

with app.app_context():
    for xl_name, db_name in MANUAL_MAP.items():
        if db_name is None:
            print(f"  SKIP: {xl_name}")
            continue
        cos = co_map.get(xl_name.upper(), co_map.get(xl_name, []))
        if not cos:
            print(f"  NO COs FOUND: {xl_name}")
            continue
        subject = Subject.query.filter_by(subject_name=db_name).first()
        if not subject:
            print(f"  DB NOT FOUND: {db_name}")
            continue
        for co_entry in cos:
            existing = CourseOutcome.query.filter_by(
                subject_id=subject.id, co_number=co_entry['co']
            ).first()
            if existing:
                skipped += 1
                continue
            db.session.add(CourseOutcome(
                subject_id=subject.id,
                co_number=co_entry['co'],
                statement=co_entry['statement']
            ))
            inserted += 1
        print(f"  OK: {xl_name} -> {db_name} ({len(cos)} COs)")
    db.session.commit()

print(f"\nInserted: {inserted}  Skipped: {skipped}")